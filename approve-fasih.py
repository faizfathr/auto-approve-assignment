import json
import random
import re
import sys
from contextlib import suppress
from decimal import Decimal, InvalidOperation
from math import ceil
from pathlib import Path
from time import monotonic
from urllib.parse import quote, urlsplit
from camoufox.fingerprints import from_fpgen, generate_fingerprint
from camoufox.sync_api import Camoufox
from camoufox.utils import get_screen_cons
from openpyxl import load_workbook
from playwright.sync_api import Error, expect

TIMEOUT = 30000

PESAN_BOT = re.compile(
    r"\bBOT-\d+\b|\bbot\s+detected\b",
    re.I
)

SELECTORS = {
    "approve": 'button:visible:has(svg.tabler-icon-check)',
    "konfirmasi_approve": '[role="alertdialog"]:visible:has-text("Approve Assignment?") button:text-is("Konfirmasi")',
    "loading": '[aria-busy="true"]:visible, [role="progressbar"]:visible, .animate-spin:visible, [class~="tw:animate-spin"]:visible, [class~="tabler-icon-loader-2"]:visible',
    "modal": r'div.tw\:border\!.tw\:relative.tw\:flex.tw\:w-full.tw\:flex-col.tw\:gap-4.tw\:bg-background.tw\:p-6.tw\:pt-4.tw\:shadow-lg\!.tw\:sm\:rounded-lg:visible',
}


def identitas(path):
    if path.exists():
        return json.loads(path.read_text(encoding="utf-8"))
    fp = generate_fingerprint(
        os="windows",
        screen=get_screen_cons()
    )
    fp.setdefault("window", {})["screenY"] = (
        from_fpgen(fp)["window.screenY"]
    )
    path.parent.mkdir(exist_ok=True)
    path.write_text(
        json.dumps(fp),
        encoding="utf-8"
    )
    return fp

def sesi(browser, page, path, simpan=False):
    if simpan:
        tab = json.loads(path.read_text(encoding="utf-8"))["tab"] if path.exists() else {}
        if not page.is_closed() and page.url.startswith("https://fasih-sm.bps.go.id/"):
            tab = page.evaluate("() => ({...sessionStorage})")
        data = {"browser": browser.storage_state(indexed_db=True), "tab": tab}
        temp = path.with_suffix(".tmp.json")
        temp.write_text(json.dumps(data), encoding="utf-8")
        temp.replace(path)
    elif path.exists():
        data = json.loads(path.read_text(encoding="utf-8"))
        browser.set_storage_state(data["browser"])
        page.add_init_script("if (location.origin === 'https://fasih-sm.bps.go.id' && !sessionStorage.length) "
                             "for (const [k,v] of Object.entries(" + json.dumps(data["tab"]) + ")) sessionStorage.setItem(k,v);")

def sukses_kirim(page):
    return (
        page.get_by_text(re.compile(r"Assignment\s+Berhasil\s+Disubmit",re.I))
        .or_(
            page.get_by_role(
                "button",
                name="Tetap di Halaman",
                exact=True
            )
        )
        .first
    )

def simpan(book, path):
    temp = path.with_suffix(".tmp.xlsx")
    book.save(temp)
    temp.replace(path)

def kirim(page):
    while not sukses_kirim(page).is_visible():
        konfirmasi = page.get_by_role(
            "button",
            name="Konfirmasi",
            exact=True
        )
        kirim_btn = page.get_by_role(
            "button",
            name="Kirim",
            exact=True
        )
        if konfirmasi.is_visible():
            konfirmasi.click()
            page.wait_for_timeout(TIMEOUT)
            continue
        if kirim_btn.is_visible():
            kirim_btn.click()
            page.wait_for_timeout(TIMEOUT)
            continue
        page.wait_for_timeout(TIMEOUT)
    return True

def gangguan(page):
    bot = page.get_by_text(PESAN_BOT).and_(page.locator(':not(#se2026_nested [data-nested-view="true"], #se2026_nested [data-nested-view="true"] *, [id^="nama_usaha_edit#"], [id^="nama_usaha_edit#"] *)'))
    return bot.filter(visible=True).first

def posisi(page):
    if sukses_kirim(page).is_visible():
        return "sukses"
    if page.locator(SELECTORS["loading"]).count():
        return "memuat"
    # if tampilan_kirim(page).count():
    #     return "pengiriman"
    if page.get_by_role("button", name="Review", exact=True).filter(visible=True).count():
        return "detail assignment"
    if page.locator(SELECTORS["approve"]).is_visible():
        return "approve"
    kirim_awal = page.get_by_role("button", name="Kirim", exact=True).filter(visible=True).count()
    if kirim_awal:
        return "kirim awal"
    if page.get_by_role("button", name="Ringkasan", exact=True).filter(visible=True).count():
        return "ringkasan"
    # if halaman_awal(page):
    #     return "survei/beranda"
    return "memuat"

def posisi_stabil(page):
    # ponytail: dua pembacaan menyaring glitch singkat; gunakan sinyal render selesai jika tersedia.
    sebelum = posisi(page)
    akhir = monotonic() + TIMEOUT / 1000
    while monotonic() < akhir:
        page.wait_for_timeout(100)
        sekarang = posisi(page)
        if sekarang == sebelum:
            return sekarang
        sebelum = sekarang
    raise Error("Tampilan masih berubah; periksa ulang posisi halaman.")

def posisi_stabil(page):
    # ponytail: dua pembacaan menyaring glitch singkat; gunakan sinyal render selesai jika tersedia.
    sebelum = posisi(page)
    akhir = monotonic() + TIMEOUT / 1000
    while monotonic() < akhir:
        page.wait_for_timeout(100)
        sekarang = posisi(page)
        if sekarang == sebelum:
            return sekarang
        sebelum = sekarang
    raise Error("Tampilan masih berubah; periksa ulang posisi halaman.")

def pasang_refresh(page, pulih=None):
    pesan = gangguan(page)
    refresh = page.get_by_role("button", name="Refresh Page", exact=True).filter(visible=True)
    def ulang():
        print("  [BOT] Hard refresh di tab yang sama.", flush=True)
        if pulih is not None:
            pulih()
        page.wait_for_timeout(1000)
        try:
            with page.expect_navigation(wait_until="domcontentloaded", timeout=TIMEOUT):
                # ponytail: reload(true) khusus Firefox/Camoufox; gunakan API bypass-cache bila pindah browser.
                page.evaluate("() => location.reload(true)")
        except Error as error:
            if page.is_closed() or re.search(r"NS_ERROR_UNKNOWN_HOST|ERR_NAME_NOT_RESOLVED", str(error)):
                raise
    page.remove_locator_handler(pesan)
    page.add_locator_handler(pesan, ulang, no_wait_after=True)
    page.remove_locator_handler(refresh)
    def refresh_web():
        if refresh.evaluate_all("buttons => { const b = buttons.find(b => !b.disabled); if (!b) return false; b.click(); return true; }"):
            print("  [REFRESH] Tombol web diklik; memeriksa ulang halaman.", flush=True)
            if pulih is not None:
                pulih()
    page.add_locator_handler(refresh, refresh_web, no_wait_after=True)

def tab_survei(browser):
    browser.cookies()
    pages = [p for p in reversed(browser.pages) if not p.is_closed() and p.url.startswith("https://fasih-sm.bps.go.id/")]
    return next((p for p in pages if urlsplit(p.url).path == "/app/surveys"), next(iter(pages), None))

def ulang_langkah(page, tahap, aksi, pulih):
    semula = posisi(page)
    while True:
        if pulih():
            raise Error("Halaman direfresh; ulangi pengisian dari awal.")
        try:
            hasil = aksi()
            if pulih():
                raise Error("Halaman direfresh; ulangi pengisian dari awal.")
            return hasil
        except (Error, AssertionError, ValueError) as error:
            if pulih() or page.is_closed() or re.search(r"NS_ERROR_UNKNOWN_HOST|ERR_NAME_NOT_RESOLVED", str(error)):
                raise
            if isinstance(error, AssertionError) and "timeout" not in str(error).lower():
                raise
            sekarang = posisi_stabil(page)
            if sekarang != semula:
                raise Error(f"Tampilan berubah dari {semula} ke {sekarang}; tentukan ulang aksi.") from error
            if sekarang == "ringkasan":
                raise Error("Akses edit tidak tersedia; buka ulang Review dari detail assignment.") from error
            if isinstance(error, ValueError):
                raise
            print("  [TUNGGU] " + tahap + " | " + str(error).splitlines()[0][:160], flush=True)
            page.wait_for_timeout(100)
def halaman_awal(page):
    return page.url.startswith("https://fasih-sm.bps.go.id/") and urlsplit(page.url).path.rstrip("/") in ("", "/app", "/app/home", "/app/dashboard", "/app/beranda", "/app/surveys")

def jalankan(pemulihan=False):
    root = Path(__file__).resolve().parent
    cfg = json.loads((root / "config.json").read_text(encoding="utf-8"))
    for key in ("delay", "batch", "rest"):
        low, high = cfg[key + "_min"], cfg[key + "_max"]
        if type(low) not in (int, float) or type(high) not in (int, float) or not 0 <= low <= high < float("inf"):
            raise ValueError("Rentang config tidak valid: " + key)
    if any(type(cfg[k]) is not int or cfg[k] < 1 for k in ("batch_min", "batch_max")):
        raise ValueError("Batch harus bilangan bulat positif")
    path = root / cfg["excel_file"]
    book = load_workbook(path)
    sheet = book[cfg["sheet_name"]]
    headers = [str(c.value).strip() for c in sheet[1]]
    headers = ["remarks" if h == "remark" else h for h in headers]
    assert all(headers.count(k) == 1 for k in ("id_assignment", "nama_kk"))
    assert all(headers.count(k) <= 1 for k in ("status", "catatan"))
    rows = []
    for number, cells in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        row = dict(zip(headers, cells))
        try:
            if (
                not isinstance(row["id_assignment"], str)
                or not row["id_assignment"].strip()
            ):
                raise ValueError("ID Assignment wajib berupa teks tidak kosong")

        except (ValueError, InvalidOperation) as error:
            if "--cek-excel" in sys.argv:
                raise
            print(error)
            continue
        rows.append((number, row))
    status = headers.index("status") + 1 if "status" in headers else sheet.max_column + 1
    sheet.cell(1, status, "status")
    keterangan = headers.index("keterangan") + 1 if "keterangan" in headers else sheet.max_column + 1
    sheet.cell(1, keterangan, "keterangan")
    catatan = headers.index("catatan") + 1 if "catatan" in headers else sheet.max_column + 1
    sheet.cell(1, catatan, "catatan")
    with path.open("r+b"):
        pass
    fingerprint = identitas(root / ".fasih-profile" / "fingerprint.json")
    headless = False
    print(f"[MULAI] {len(rows)} baris; headless={headless}. Membuka browser dan memuat sesi tersimpan.", flush=True)
    with Camoufox(persistent_context=True, user_data_dir=str(root / ".fasih-profile"), fingerprint=fingerprint, i_know_what_im_doing=True, headless=headless, humanize=False, os="windows", locale="id-ID", timeout=TIMEOUT) as browser:
        browser.set_default_timeout(TIMEOUT)
        browser.set_default_navigation_timeout(TIMEOUT)
        dimuat_ulang = False
        def pulih_bot():
            nonlocal dimuat_ulang
            dimuat_ulang = True
        def tab_baru(tab):
            try:
                if sum(not p.is_closed() for p in browser.pages) > 2:
                    tab.close()
                elif not tab.is_closed():
                    pasang_refresh(tab, pulih_bot)
            except Error:
                if not tab.is_closed():
                    raise
        browser.on("page", tab_baru)
        for tab in browser.pages:
            pasang_refresh(tab, pulih_bot)
        page = browser.pages[0] if browser.pages else browser.new_page()
        for tab in browser.pages:
            if tab != page:
                tab.close()
        sesi(browser, page, root / ".fasih-profile" / "sesi.json")
        print("[SESI] Memeriksa akses halaman Survei...", flush=True)
        dibuka = False
        while True:
            try:
                if not dibuka:
                    dibuka = True
                    page.goto("https://fasih-sm.bps.go.id/app/surveys?page=0&perPage=10&layout=list", wait_until="commit")
                expect(page.locator("body")).to_be_visible(timeout=TIMEOUT)
                expect(gangguan(page)).to_be_hidden(timeout=TIMEOUT)
                if not pemulihan:
                    pemulihan = True
                    input("Setelah login awal selesai, tekan Enter sekali untuk mulai otomatis... ")
                tab = tab_survei(browser)
                if tab is None:
                    page.wait_for_timeout(100)
                    continue
                page = tab
                if urlsplit(page.url).path != "/app/surveys":
                    page.goto("https://fasih-sm.bps.go.id/app/surveys?page=0&perPage=10&layout=list", wait_until="commit")
                page.wait_for_url("**/app/surveys?*", wait_until="commit")
                sesi(browser, page, root / ".fasih-profile" / "sesi.json", simpan=True)
                break
            except (Error, AssertionError, StopIteration) as error:
                if re.search(r"NS_ERROR_UNKNOWN_HOST|ERR_NAME_NOT_RESOLVED", str(error)):
                    raise
                if page.is_closed():
                    raise
                page.wait_for_timeout(100)
        print("[SESI] Halaman Survei terbuka. Mulai memeriksa assignment.", flush=True)
        hasil = dict.fromkeys(("SUCCESS", "NO_CHANGE", "ERROR"), 0)
        dimuat_ulang = False
        sisa = random.randint(cfg["batch_min"], cfg["batch_max"])
        for index, (number, row) in enumerate(rows):
            mulai_baris = monotonic()
            print(f"[{index + 1}/{len(rows)}] Baris Excel {number}: {row['id_assignment']}", flush=True)
            sheet.cell(number, status, "ERROR")
            simpan(book, path)
            assignment = page
            result, note = "ERROR", ""
            kirim_ulang = "Kirim belum selesai" in str(row.get("catatan") or "")
            tahap = "Membuka assignment detail"
            print("  " + tahap, flush=True)
            detail_dibuka = False
            isi_ulang = False
            while True:
                try:
                    if dimuat_ulang:
                        dimuat_ulang = False
                        isi_ulang = True
                        print("  [PULIH] Mengenali ulang tampilan setelah refresh.", flush=True)
                    if (detail_dibuka or page != assignment) and halaman_awal(page):
                        assignment = page
                        detail_dibuka = False
                        isi_ulang = True
                        print("  [PULIH] Survei/Beranda terbuka; membuka kembali assignment yang sama.", flush=True)
                    if not detail_dibuka:
                        detail_dibuka = True
                        for tab in browser.pages:
                            if tab != assignment:
                                tab.close()
                        pasang_refresh(page, pulih_bot)
                        page.goto("https://fasih-sm.bps.go.id/app/assignment-detail/" + quote(row["id_assignment"], safe=""), wait_until="commit")
                    review = page.get_by_role("button", name="Review", exact=True).filter(visible=True)
                    with page.expect_popup(timeout=TIMEOUT) as popup:
                        review.click()
                    form = popup.value
                    tahap = "Memeriksa tampilan form/pengiriman"
                    print("  " + tahap, flush=True)
                    approve = form.locator(SELECTORS["approve"]).filter(visible=True)
                    try:
                        expect(approve).to_be_visible(timeout=TIMEOUT)
                        print("  " + tahap + " | Tombol Approve tersedia.", flush=True)
                        approve.first.click()
                        konfirmasi_approve = form.locator(
                            SELECTORS["konfirmasi_approve"]
                        )
                        expect(konfirmasi_approve).to_be_visible(
                            timeout=TIMEOUT
                        )
                        konfirmasi_approve.click()
                        kirim(form)
                        result = "SUCCESS"
                        sheet.cell(number, status, result)
                        break
                    except AssertionError:
                        result = "NO_CHANGE"
                        note = "Tidak ada tombol Approve atau assignment sudah diproses"
                        print(f"  [SKIP] {note}", flush=True)
                        result = "NO_CHANGE"
                        pass
                        sheet.cell(number, status, result)
                    if index + 1 < len(rows):
                        jeda = random.randint(5, 15)
                        print(
                            f"  [JEDA] {jeda} detik...",
                            flush=True
                        )
                        page.wait_for_timeout(jeda * 1000)
                except (Error, AssertionError, ValueError, InvalidOperation) as error:
                    error_text = str(error).strip()
                    if re.search(r"NS_ERROR_UNKNOWN_HOST|ERR_NAME_NOT_RESOLVED", error_text):
                        sheet.cell(number, status, "RETRY")
                        simpan(book, path)
                        raise
                    if result != "ERROR":
                        break
                    # if halaman_awal(page):
                    #     continue
                    if dimuat_ulang:
                        if page.is_closed():
                            raise
                        continue
                    if isinstance(error, (ValueError, InvalidOperation)) or (isinstance(error, AssertionError) and "timeout" not in error_text.lower()):
                        note = f"{tahap} | {type(error).__name__}: {error_text}"
                        break
                    if page.is_closed():
                        raise
                    sekarang = posisi_stabil(page)
                    if sekarang == "ringkasan" or (tahap in ("Menunggu Revoke/Edit", "Revoke dan konfirmasi", "Membuka Edit") and sekarang in ("approve", "konfirmasi revoke")):
                        print("  [PULIH AKSES] Kembali ke detail assignment yang sama; membuka Review lagi.", flush=True)
                        if page != assignment:
                            page.close()
                        if assignment.is_closed():
                            assignment = browser.new_page()
                        page = assignment
                        page.bring_to_front()
                        detail_dibuka = urlsplit(page.url).path == "/app/assignment-detail/" + quote(row["id_assignment"], safe="")
                        isi_ulang = True
                        continue
                    print("  [TUNGGU] " + tahap + " | Posisi: " + sekarang + " | " + error_text.splitlines()[0][:160], flush=True)
                    page.wait_for_timeout(100)
if __name__ == "__main__":
    jalankan()